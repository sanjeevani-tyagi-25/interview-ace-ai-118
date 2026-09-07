import io
import json
import os
import time

from dotenv import load_dotenv
from google import genai
from docx import Document
from pypdf import PdfReader

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
)


# ============================================================
# LOAD ENVIRONMENT VARIABLES
# ============================================================

load_dotenv()


# ============================================================
# GEMINI CONFIGURATION
# ============================================================

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

if not GEMINI_API_KEY:
    print(
        "⚠️ GEMINI_API_KEY is not configured. "
        "Resume analysis will use fallback mode."
    )
    client = None
else:
    client = genai.Client(
        api_key=GEMINI_API_KEY
    )


# ============================================================
# GEMINI MODEL
# ============================================================

GEMINI_MODEL = os.getenv(
    "GEMINI_RESUME_MODEL",
    "gemini-3.5-flash"
)


# ============================================================
# EXTRACT PDF TEXT
# ============================================================

def extract_pdf_text(content: bytes) -> str:

    try:

        reader = PdfReader(
            io.BytesIO(content)
        )

        text = []

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                text.append(page_text)

        return "\n".join(text)

    except Exception as error:

        print(
            "❌ PDF extraction error:",
            error
        )

        raise ValueError(
            "Could not read the PDF file."
        )


# ============================================================
# EXTRACT DOCX TEXT
# ============================================================

def extract_docx_text(content: bytes) -> str:

    try:

        document = Document(
            io.BytesIO(content)
        )

        text = []

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                text.append(
                    paragraph.text
                )

        # Also extract text from tables
        for table in document.tables:

            for row in table.rows:

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    if cell_text:

                        text.append(
                            cell_text
                        )

        return "\n".join(text)

    except Exception as error:

        print(
            "❌ DOCX extraction error:",
            error
        )

        raise ValueError(
            "Could not read the DOCX file."
        )


# ============================================================
# EXTRACT RESUME TEXT
# ============================================================

def extract_text(
    filename: str,
    content: bytes,
) -> str:

    filename = filename.lower()

    if filename.endswith(".pdf"):

        return extract_pdf_text(
            content
        )

    if filename.endswith(".docx"):

        return extract_docx_text(
            content
        )

    raise ValueError(
        "Only PDF and DOCX files are supported."
    )


# ============================================================
# CLEAN GEMINI JSON RESPONSE
# ============================================================

def clean_json_response(text: str) -> str:

    text = text.strip()

    if text.startswith("```json"):

        text = text[
            len("```json"):
        ]

    elif text.startswith("```"):

        text = text[
            len("```"):
        ]

    if text.endswith("```"):

        text = text[
            :-len("```")
        ]

    text = text.strip()

    start = text.find("{")
    end = text.rfind("}")

    if start == -1 or end == -1:

        raise ValueError(
            "Gemini response did not contain a valid JSON object."
        )

    return text[
        start:end + 1
    ]


# ============================================================
# FALLBACK RESUME ANALYSIS
# ============================================================

def fallback_resume_analysis(
    resume_text: str,
) -> dict:

    """
    Local fallback used when Gemini is unavailable.

    This does not pretend to be AI analysis.
    It simply checks the extracted resume text for
    common resume sections and technical keywords.
    """

    text = resume_text.lower()

    # --------------------------------------------------------
    # COMMON TECHNICAL SKILLS
    # --------------------------------------------------------

    possible_skills = [
        "python",
        "java",
        "c++",
        "c",
        "javascript",
        "typescript",
        "react",
        "node.js",
        "node",
        "html",
        "css",
        "sql",
        "mongodb",
        "mysql",
        "postgresql",
        "fastapi",
        "flask",
        "django",
        "tensorflow",
        "keras",
        "pytorch",
        "scikit-learn",
        "machine learning",
        "deep learning",
        "data science",
        "git",
        "github",
        "docker",
        "aws",
        "azure",
        "power bi",
        "pandas",
        "numpy",
    ]

    skills_found = []

    for skill in possible_skills:

        if skill in text:

            # Keep original-style capitalization
            display_skill = skill

            replacements = {
                "python": "Python",
                "java": "Java",
                "c++": "C++",
                "c": "C",
                "javascript": "JavaScript",
                "typescript": "TypeScript",
                "react": "React",
                "node.js": "Node.js",
                "node": "Node.js",
                "html": "HTML",
                "css": "CSS",
                "sql": "SQL",
                "mongodb": "MongoDB",
                "mysql": "MySQL",
                "postgresql": "PostgreSQL",
                "fastapi": "FastAPI",
                "flask": "Flask",
                "django": "Django",
                "tensorflow": "TensorFlow",
                "keras": "Keras",
                "pytorch": "PyTorch",
                "scikit-learn": "Scikit-learn",
                "machine learning": "Machine Learning",
                "deep learning": "Deep Learning",
                "data science": "Data Science",
                "git": "Git",
                "github": "GitHub",
                "docker": "Docker",
                "aws": "AWS",
                "azure": "Azure",
                "power bi": "Power BI",
                "pandas": "Pandas",
                "numpy": "NumPy",
            }

            display_skill = replacements.get(
                skill,
                skill
            )

            if display_skill not in skills_found:

                skills_found.append(
                    display_skill
                )

    # --------------------------------------------------------
    # SECTION CHECKS
    # --------------------------------------------------------

    sections = {
        "education": any(
            word in text
            for word in [
                "education",
                "b.tech",
                "btech",
                "bachelor",
                "degree",
                "university",
                "college",
            ]
        ),

        "experience": any(
            word in text
            for word in [
                "experience",
                "internship",
                "intern",
                "employment",
                "worked",
            ]
        ),

        "projects": any(
            word in text
            for word in [
                "projects",
                "project",
            ]
        ),

        "achievements": any(
            word in text
            for word in [
                "achievement",
                "award",
                "winner",
                "certification",
                "certificate",
            ]
        ),

        "contact": any(
            word in text
            for word in [
                "@",
                "linkedin",
                "github",
                "phone",
            ]
        ),
    }

    # --------------------------------------------------------
    # SCORE
    # --------------------------------------------------------

    score = 40

    if sections["contact"]:
        score += 10

    if sections["education"]:
        score += 10

    if sections["experience"]:
        score += 10

    if sections["projects"]:
        score += 10

    if sections["achievements"]:
        score += 5

    if len(skills_found) >= 5:
        score += 10

    score = min(
        score,
        95
    )

    # --------------------------------------------------------
    # MISSING SKILLS
    # --------------------------------------------------------

    missing_skills = []

    recommended_skills = [
        "Git",
        "SQL",
        "Problem Solving",
        "Communication",
    ]

    for skill in recommended_skills:

        if skill.lower() not in text:

            missing_skills.append(
                skill
            )

    # --------------------------------------------------------
    # STRENGTHS
    # --------------------------------------------------------

    strengths = []

    if skills_found:

        strengths.append(
            f"Technical skills identified: "
            f"{', '.join(skills_found[:8])}."
        )

    if sections["projects"]:

        strengths.append(
            "Projects section is present."
        )

    if sections["education"]:

        strengths.append(
            "Education information is present."
        )

    if sections["experience"]:

        strengths.append(
            "Experience or internship information is present."
        )

    if not strengths:

        strengths.append(
            "Resume text was successfully extracted."
        )

    # --------------------------------------------------------
    # WEAKNESSES
    # --------------------------------------------------------

    weaknesses = []

    if not sections["experience"]:

        weaknesses.append(
            "No clear experience or internship section was detected."
        )

    if not sections["projects"]:

        weaknesses.append(
            "No clear projects section was detected."
        )

    if not sections["achievements"]:

        weaknesses.append(
            "No clear achievements or certifications section was detected."
        )

    if not skills_found:

        weaknesses.append(
            "No common technical skills were detected."
        )

    # --------------------------------------------------------
    # SUGGESTIONS
    # --------------------------------------------------------

    suggestions = [
        "Add measurable results to projects and experience.",
        "Use job-specific technical keywords relevant to the target role.",
        "Keep formatting simple and ATS-friendly.",
    ]

    if not sections["experience"]:

        suggestions.append(
            "Add internship, training, or relevant practical experience."
        )

    if not sections["projects"]:

        suggestions.append(
            "Add 2–3 relevant technical projects with technologies used."
        )

    # --------------------------------------------------------
    # RESULT
    # --------------------------------------------------------

    return {
        "atsScore": score,

        "skillsFound": skills_found,

        "missingSkills": missing_skills,

        "suggestions": suggestions,

        "strengths": strengths,

        "weaknesses": weaknesses,

        "summary": (
            "The resume was successfully processed locally. "
            "AI analysis was temporarily unavailable, so this "
            "result is based on local resume text and section checks."
        ),

        "experienceFeedback": (
            "Local analysis completed. "
            "For detailed AI feedback, Gemini analysis can be "
            "performed when the service is available."
        ),

        "projectsFeedback": (
            "Local analysis completed. "
            "Projects were checked based on extracted resume text."
        ),

        "keywordFeedback": (
            f"{len(skills_found)} common technical skills "
            "were detected in the resume."
        ),
    }


# ============================================================
# AI RESUME ANALYSIS
# ============================================================

def analyze_with_ai(
    resume_text: str,
):

    prompt = f"""
You are an expert ATS resume analyzer and senior technical recruiter.

Analyze the following resume carefully.

RESUME:
{resume_text}

Return ONLY valid JSON.

Do not include:
- Markdown
- ```json
- ```
- Explanations outside the JSON
- Any text before or after the JSON

Use exactly this structure:

{{
    "atsScore": 0,
    "skillsFound": [],
    "missingSkills": [],
    "suggestions": [],
    "strengths": [],
    "weaknesses": [],
    "summary": "",
    "experienceFeedback": "",
    "projectsFeedback": "",
    "keywordFeedback": ""
}}

Rules:

1. atsScore must be an integer from 0 to 100.
2. Do NOT automatically give a high score.
3. Judge the resume based on:
   - skills
   - projects
   - education
   - experience
   - achievements
   - formatting/readability
   - relevant keywords
   - measurable results
   - technical relevance
4. skillsFound must contain technical and professional skills actually present in the resume.
5. missingSkills should contain useful skills or keywords that appear absent or weak.
6. suggestions must contain specific actionable improvements.
7. strengths must contain real strengths found in the resume.
8. weaknesses must contain actual weaknesses.
9. summary should briefly describe the candidate's resume quality.
10. experienceFeedback should evaluate the experience section.
11. projectsFeedback should evaluate the projects section.
12. keywordFeedback should explain how well the resume uses job-relevant keywords.
13. Do not invent:
   - experience
   - skills
   - education
   - companies
   - achievements
14. Keep suggestions concise and useful.
15. All strings must use valid JSON formatting.
16. Arrays must contain strings only.
17. atsScore must be a number, not a string.
"""

    if client is None:

        print(
            "⚠️ Gemini client unavailable."
        )

        return fallback_resume_analysis(
            resume_text
        )

    max_retries = 3

    for attempt in range(max_retries):

        try:

            print(
                "\n========================================"
            )

            print(
                "🤖 Sending resume to Gemini..."
            )

            print(
                f"Model: {GEMINI_MODEL}"
            )

            print(
                f"Attempt: {attempt + 1}/{max_retries}"
            )

            print(
                "========================================"
            )

            response = client.models.generate_content(
                model=GEMINI_MODEL,
                contents=prompt,
            )

            if not response:

                raise ValueError(
                    "Gemini returned no response."
                )

            if not response.text:

                raise ValueError(
                    "Gemini returned an empty response."
                )

            text = response.text.strip()

            print(
                "\n========== RAW GEMINI RESPONSE =========="
            )

            print(text)

            print(
                "=========================================="
            )

            json_text = clean_json_response(
                text
            )

            result = json.loads(
                json_text
            )

            required_fields = [
                "atsScore",
                "skillsFound",
                "missingSkills",
                "suggestions",
                "strengths",
                "weaknesses",
                "summary",
                "experienceFeedback",
                "projectsFeedback",
                "keywordFeedback",
            ]

            for field in required_fields:

                if field not in result:

                    raise ValueError(
                        f"Gemini response is missing field: {field}"
                    )

            try:

                result["atsScore"] = int(
                    result["atsScore"]
                )

            except Exception:

                raise ValueError(
                    "Invalid ATS score returned by Gemini."
                )

            if not 0 <= result["atsScore"] <= 100:

                raise ValueError(
                    "ATS score must be between 0 and 100."
                )

            array_fields = [
                "skillsFound",
                "missingSkills",
                "suggestions",
                "strengths",
                "weaknesses",
            ]

            for field in array_fields:

                if not isinstance(
                    result[field],
                    list
                ):

                    raise ValueError(
                        f"{field} must be an array."
                    )

            print(
                "\n✅ Resume JSON parsed successfully."
            )

            print(
                f"📊 ATS Score: "
                f"{result['atsScore']}/100"
            )

            return result

        except Exception as error:

            error_text = str(error)

            print(
                "\n❌ Gemini resume analysis failed:"
            )

            print(
                error_text
            )

            is_temporary_error = (
                "503" in error_text
                or "UNAVAILABLE" in error_text
                or "high demand" in error_text.lower()
                or "temporarily unavailable" in error_text.lower()
                or "429" in error_text
                or "RESOURCE_EXHAUSTED" in error_text
                or "quota" in error_text.lower()
            )

            if (
                is_temporary_error
                and attempt < max_retries - 1
            ):

                wait_time = 3

                print(
                    "\n⏳ Gemini is temporarily unavailable."
                )

                print(
                    f"Retrying in {wait_time} seconds..."
                )

                time.sleep(
                    wait_time
                )

                continue

            # ------------------------------------------------
            # IMPORTANT:
            # DO NOT CRASH RESUME UPLOAD
            # ------------------------------------------------

            print(
                "\n⚠️ Gemini could not analyze the resume."
            )

            print(
                "➡️ Switching to local fallback analysis."
            )

            return fallback_resume_analysis(
                resume_text
            )


# ============================================================
# MAIN RESUME ANALYZER
# ============================================================

def analyze_resume(
    filename: str,
    content: bytes,
):

    print(
        "\n========================================"
    )

    print(
        "📄 STARTING RESUME ANALYSIS"
    )

    print(
        "========================================"
    )

    # --------------------------------------------------------
    # EXTRACT TEXT
    # --------------------------------------------------------

    resume_text = extract_text(
        filename,
        content,
    )

    print(
        f"📄 Extracted resume text length: "
        f"{len(resume_text)} characters"
    )

    # --------------------------------------------------------
    # CHECK TEXT
    # --------------------------------------------------------

    if not resume_text.strip():

        raise ValueError(
            "Could not extract any text from the resume."
        )

    # --------------------------------------------------------
    # ANALYZE
    # --------------------------------------------------------

    result = analyze_with_ai(
        resume_text
    )

    # --------------------------------------------------------
    # PRINT RESULT
    # --------------------------------------------------------

    print(
        "\n========== FINAL RESUME RESULT =========="
    )

    print(
        json.dumps(
            result,
            indent=2,
            ensure_ascii=False,
        )
    )

    print(
        "=========================================="
    )

    return result


# ============================================================
# GENERATE PDF REPORT
# ============================================================

def generate_resume_report(
    result: dict,
) -> bytes:

    buffer = io.BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]

    title_style.alignment = TA_CENTER

    heading_style = styles["Heading2"]

    body_style = styles["BodyText"]

    body_style.leading = 16

    story = []

    # ========================================================
    # TITLE
    # ========================================================

    story.append(
        Paragraph(
            "InterviewAI - Resume Analysis Report",
            title_style,
        )
    )

    story.append(
        Spacer(1, 20)
    )

    # ========================================================
    # ATS SCORE
    # ========================================================

    score = result.get(
        "atsScore",
        0,
    )

    story.append(
        Paragraph(
            f"<b>ATS Score: {score}/100</b>",
            heading_style,
        )
    )

    story.append(
        Spacer(1, 12)
    )

    # ========================================================
    # SUMMARY
    # ========================================================

    story.append(
        Paragraph(
            "AI Resume Summary",
            heading_style,
        )
    )

    summary = result.get(
        "summary",
        "",
    )

    story.append(
        Paragraph(
            str(summary),
            body_style,
        )
    )

    story.append(
        Spacer(1, 15)
    )

    # ========================================================
    # SKILLS FOUND
    # ========================================================

    story.append(
        Paragraph(
            "Skills Found",
            heading_style,
        )
    )

    skills = result.get(
        "skillsFound",
        [],
    )

    if skills:

        story.append(
            Paragraph(
                ", ".join(
                    str(skill)
                    for skill in skills
                ),
                body_style,
            )
        )

    else:

        story.append(
            Paragraph(
                "No specific skills were identified.",
                body_style,
            )
        )

    story.append(
        Spacer(1, 15)
    )

    # ========================================================
    # MISSING SKILLS
    # ========================================================

    story.append(
        Paragraph(
            "Missing Skills",
            heading_style,
        )
    )

    missing = result.get(
        "missingSkills",
        [],
    )

    if missing:

        story.append(
            Paragraph(
                ", ".join(
                    str(skill)
                    for skill in missing
                ),
                body_style,
            )
        )

    else:

        story.append(
            Paragraph(
                "No major missing skills identified.",
                body_style,
            )
        )

    story.append(
        Spacer(1, 15)
    )

    # ========================================================
    # STRENGTHS
    # ========================================================

    story.append(
        Paragraph(
            "Strengths",
            heading_style,
        )
    )

    strengths = result.get(
        "strengths",
        [],
    )

    if strengths:

        for item in strengths:

            story.append(
                Paragraph(
                    f"• {item}",
                    body_style,
                )
            )

    else:

        story.append(
            Paragraph(
                "No strengths were identified.",
                body_style,
            )
        )

    story.append(
        Spacer(1, 15)
    )

    # ========================================================
    # WEAKNESSES
    # ========================================================

    story.append(
        Paragraph(
            "Weaknesses",
            heading_style,
        )
    )

    weaknesses = result.get(
        "weaknesses",
        [],
    )

    if weaknesses:

        for item in weaknesses:

            story.append(
                Paragraph(
                    f"• {item}",
                    body_style,
                )
            )

    else:

        story.append(
            Paragraph(
                "No major weaknesses were identified.",
                body_style,
            )
        )

    story.append(
        Spacer(1, 15)
    )

    # ========================================================
    # SUGGESTIONS
    # ========================================================

    story.append(
        Paragraph(
            "AI Suggestions",
            heading_style,
        )
    )

    suggestions = result.get(
        "suggestions",
        [],
    )

    if suggestions:

        for item in suggestions:

            story.append(
                Paragraph(
                    f"• {item}",
                    body_style,
                )
            )

    else:

        story.append(
            Paragraph(
                "No suggestions available.",
                body_style,
            )
        )

    story.append(
        Spacer(1, 15)
    )

    # ========================================================
    # EXPERIENCE FEEDBACK
    # ========================================================

    story.append(
        Paragraph(
            "Experience Feedback",
            heading_style,
        )
    )

    experience_feedback = result.get(
        "experienceFeedback",
        "",
    )

    story.append(
        Paragraph(
            str(experience_feedback),
            body_style,
        )
    )

    story.append(
        Spacer(1, 15)
    )

    # ========================================================
    # PROJECTS FEEDBACK
    # ========================================================

    story.append(
        Paragraph(
            "Projects Feedback",
            heading_style,
        )
    )

    projects_feedback = result.get(
        "projectsFeedback",
        "",
    )

    story.append(
        Paragraph(
            str(projects_feedback),
            body_style,
        )
    )

    story.append(
        Spacer(1, 15)
    )

    # ========================================================
    # KEYWORD FEEDBACK
    # ========================================================

    story.append(
        Paragraph(
            "Keyword Feedback",
            heading_style,
        )
    )

    keyword_feedback = result.get(
        "keywordFeedback",
        "",
    )

    story.append(
        Paragraph(
            str(keyword_feedback),
            body_style,
        )
    )

    # ========================================================
    # BUILD PDF
    # ========================================================

    document.build(
        story
    )

    buffer.seek(0)

    pdf_bytes = buffer.getvalue()

    print(
        f"\n✅ Resume PDF report generated: "
        f"{len(pdf_bytes)} bytes"
    )

    return pdf_bytes